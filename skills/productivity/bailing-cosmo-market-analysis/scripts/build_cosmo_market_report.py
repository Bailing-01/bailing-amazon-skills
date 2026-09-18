#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
COSMO 市场分析报告生成器
从 keyword-cosmo-attribute-*-v41.csv 产出 HTML + Excel（可筛选词明细 + 各维副表汇总）。
多标签（a|b）搜索量等分，口径在说明中标注。
"""

from __future__ import annotations

import argparse
import csv
import html
import json
import math
import re
import statistics
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as e:  # pragma: no cover
    raise SystemExit(
        "需要 openpyxl。请使用 /workspace/venv-ads/bin/python 或 pip install openpyxl"
    ) from e


# ---------------------------------------------------------------------------
# Column aliases (Chinese headers preferred)
# ---------------------------------------------------------------------------

ATTR_DIMS = [
    "目标人群",
    "修饰/价值",
    "购买场所",
    "季节/节日",
    "场景/活动",
    "产地/文化",
    "通用/品类",
    "成分/材质",
    "痛点/需求",
    "产品形态",
    "使用部位",
    "功效",
    "风格属性",
]

CORE_COLS = [
    "序号",
    "关键词",
    "搜索量",
    "流量等级",
    "竞品数",
    "蓝海度",
    "ABA转化份额",
    "准入难度",
    "竞品表现得分",
    "相关性等级",
    "推荐行动",
]

COL_ALIASES = {
    "序号": ["序号", "index", "id", "#"],
    "关键词": ["关键词", "keyword", "search term", "phrase"],
    "搜索量": ["搜索量", "search volume", "volume", "sv", "月搜索量"],
    "流量等级": ["流量等级", "traffic tier", "tier"],
    "竞品数": ["竞品数", "competitors", "competing products"],
    "蓝海度": ["蓝海度", "blue ocean", "competition"],
    "ABA转化份额": ["ABA转化份额", "aba", "aba share"],
    "准入难度": ["准入难度", "entry difficulty"],
    "竞品表现得分": ["竞品表现得分", "competitor score"],
    "相关性等级": ["相关性等级", "relevance"],
    "推荐行动": ["推荐行动", "action", "recommendation"],
}
for d in ATTR_DIMS:
    COL_ALIASES[d] = [d]


def _norm_header(h: str) -> str:
    return re.sub(r"\s+", "", (h or "").strip().lower())


def map_columns(fieldnames: Sequence[str]) -> Dict[str, str]:
    """Map canonical name -> actual CSV header."""
    actual = { _norm_header(f): f for f in fieldnames if f }
    mapping: Dict[str, str] = {}
    for canon, aliases in COL_ALIASES.items():
        for a in aliases:
            key = _norm_header(a)
            if key in actual:
                mapping[canon] = actual[key]
                break
        # also try exact match on original
        if canon not in mapping:
            for f in fieldnames:
                if f == canon:
                    mapping[canon] = f
                    break
    return mapping


def parse_float(v: Any) -> float:
    if v is None:
        return 0.0
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "—", "N/A", "n/a", "无数据", "null"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        m = re.search(r"[-+]?\d*\.?\d+", s)
        return float(m.group()) if m else 0.0


def parse_tags(raw: Any) -> List[str]:
    if raw is None:
        return []
    s = str(raw).strip()
    if s in ("", "-", "—", "N/A", "n/a", "无"):
        return []
    parts = re.split(r"[|｜;/；、]+", s)
    out = []
    seen = set()
    for p in parts:
        t = p.strip()
        if not t or t in ("-", "—"):
            continue
        if t not in seen:
            seen.add(t)
            out.append(t)
    return out


def percentile(sorted_vals: List[float], p: float) -> float:
    if not sorted_vals:
        return 0.0
    if len(sorted_vals) == 1:
        return sorted_vals[0]
    k = (len(sorted_vals) - 1) * (p / 100.0)
    f = math.floor(k)
    c = math.ceil(k)
    if f == c:
        return sorted_vals[int(k)]
    return sorted_vals[f] * (c - k) + sorted_vals[c] * (k - f)


# ---------------------------------------------------------------------------
# Load & analyze
# ---------------------------------------------------------------------------

def load_rows(path: Path) -> Tuple[List[Dict[str, Any]], Dict[str, str], List[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            raise SystemExit(f"CSV 无表头: {path}")
        colmap = map_columns(reader.fieldnames)
        if "关键词" not in colmap:
            raise SystemExit(f"缺少关键词列。实际表头: {reader.fieldnames}")
        raw_rows = list(reader)

    rows: List[Dict[str, Any]] = []
    for i, r in enumerate(raw_rows, 1):
        kw = (r.get(colmap["关键词"]) or "").strip()
        if not kw:
            continue
        vol = parse_float(r.get(colmap.get("搜索量", ""), 0))
        item: Dict[str, Any] = {
            "_i": i,
            "序号": (r.get(colmap.get("序号", ""), "") or str(i)),
            "关键词": kw,
            "搜索量": vol,
        }
        for c in CORE_COLS:
            if c in ("序号", "关键词", "搜索量"):
                continue
            if c in colmap:
                item[c] = (r.get(colmap[c]) or "").strip()
            else:
                item[c] = ""
        for d in ATTR_DIMS:
            raw = r.get(colmap[d], "") if d in colmap else ""
            tags = parse_tags(raw)
            item[d] = tags
            item[f"{d}__raw"] = (raw or "").strip() if isinstance(raw, str) else str(raw or "")
        rows.append(item)
    return rows, colmap, list(reader.fieldnames or [])


def fill_rate(rows: List[Dict[str, Any]], dim: str) -> float:
    if not rows:
        return 0.0
    n = sum(1 for r in rows if r.get(dim))
    return 100.0 * n / len(rows)


def dim_aggregate(
    rows: List[Dict[str, Any]], dim: str, *, equal_split: bool = True
) -> List[Dict[str, Any]]:
    """Aggregate tag -> keyword count, search volume (equal-split for multi-tags)."""
    tag_kw: Dict[str, set] = defaultdict(set)
    tag_vol: Dict[str, float] = defaultdict(float)
    tag_vols_list: Dict[str, List[float]] = defaultdict(list)

    for r in rows:
        tags = r.get(dim) or []
        if not tags:
            continue
        vol = float(r["搜索量"])
        share = vol / len(tags) if equal_split else vol
        for t in tags:
            tag_kw[t].add(r["关键词"])
            tag_vol[t] += share
            tag_vols_list[t].append(share)

    total_vol = sum(tag_vol.values()) or 1.0
    # Top keyword per tag (by full keyword volume among those carrying the tag)
    top_kw: Dict[str, Tuple[str, float]] = {}
    for r in rows:
        tags = r.get(dim) or []
        for t in tags:
            cur = top_kw.get(t)
            if cur is None or r["搜索量"] > cur[1]:
                top_kw[t] = (r["关键词"], r["搜索量"])

    out = []
    for t, kws in tag_kw.items():
        vols = tag_vols_list[t]
        med = statistics.median(vols) if vols else 0.0
        out.append({
            "维度": dim,
            "标签": t,
            "词数": len(kws),
            "搜索量": round(tag_vol[t], 2),
            "中位搜索量": round(med, 2),
            "占比": round(100.0 * tag_vol[t] / total_vol, 2),
            "Top词": top_kw.get(t, ("", 0))[0],
            "Top词搜索量": top_kw.get(t, ("", 0))[1],
        })
    out.sort(key=lambda x: (-x["搜索量"], -x["词数"], x["标签"]))
    return out


def head_tail_split(rows: List[Dict[str, Any]], head_pct: float = 4.0, tail_pct: float = 80.0):
    """Head = top head_pct% by volume; Tail = bottom tail_pct% by keyword count (lowest volume)."""
    sorted_rows = sorted(rows, key=lambda r: r["搜索量"], reverse=True)
    n = len(sorted_rows)
    if n == 0:
        return [], [], sorted_rows
    head_n = max(1, int(math.ceil(n * head_pct / 100.0)))
    tail_n = max(1, int(math.floor(n * tail_pct / 100.0)))
    head = sorted_rows[:head_n]
    # bottom by volume = reverse sorted take last tail_n
    tail = sorted_rows[-tail_n:] if tail_n < n else sorted_rows
    return head, tail, sorted_rows


def attr_distribution(subset: List[Dict[str, Any]], dim: str, top_n: int = 15) -> List[Dict[str, Any]]:
    return dim_aggregate(subset, dim)[:top_n]


def cooccurrence(
    rows: List[Dict[str, Any]], dim_a: str, dim_b: str, top_n: int = 20
) -> List[Dict[str, Any]]:
    pair_vol: Dict[Tuple[str, str], float] = defaultdict(float)
    pair_kw: Dict[Tuple[str, str], set] = defaultdict(set)
    for r in rows:
        ta = r.get(dim_a) or []
        tb = r.get(dim_b) or []
        if not ta or not tb:
            continue
        vol = float(r["搜索量"])
        denom = len(ta) * len(tb)
        share = vol / denom if denom else 0.0
        for a in ta:
            for b in tb:
                pair_vol[(a, b)] += share
                pair_kw[(a, b)].add(r["关键词"])
    items = [
        {
            "维A": dim_a,
            "标签A": a,
            "维B": dim_b,
            "标签B": b,
            "搜索量": round(v, 2),
            "词数": len(pair_kw[(a, b)]),
        }
        for (a, b), v in pair_vol.items()
    ]
    items.sort(key=lambda x: -x["搜索量"])
    return items[:top_n]


def blue_ocean_window(rows: List[Dict[str, Any]], top_n: int = 30) -> List[Dict[str, Any]]:
    """High volume + blue-ish competition."""
    blue_labels = {"蓝海", "温和", "blue", "mild", "low"}
    scored = []
    for r in rows:
        ocean = (r.get("蓝海度") or "").strip()
        vol = r["搜索量"]
        if vol <= 0:
            continue
        is_blue = any(b in ocean for b in blue_labels) or ocean in blue_labels
        # score: volume with boost for blue
        boost = 1.5 if is_blue else (1.2 if ocean in ("一般",) else 0.5)
        if not is_blue and ocean not in ("一般", "温和", "蓝海"):
            # still include high-volume mild only
            if ocean in ("激烈", "极激烈"):
                continue
        scored.append({
            "关键词": r["关键词"],
            "搜索量": vol,
            "蓝海度": ocean or "-",
            "竞品数": r.get("竞品数") or "",
            "推荐行动": r.get("推荐行动") or "",
            "score": vol * boost,
        })
    scored.sort(key=lambda x: -x["score"])
    return scored[:top_n]


def build_insights(rows: List[Dict[str, Any]], aggs: Dict[str, List], meta: Dict) -> List[str]:
    insights: List[str] = []
    n = len(rows)
    vols = sorted(r["搜索量"] for r in rows)
    total = sum(vols) or 1.0
    p50 = percentile(vols, 50)
    p90 = percentile(vols, 90)
    head, tail, sorted_rows = head_tail_split(rows, 4.0, 80.0)
    head_vol = sum(r["搜索量"] for r in head)
    head_share = 100.0 * head_vol / total
    head_n = len(head)
    tail_vol = sum(r["搜索量"] for r in tail)

    insights.append(
        f"共 {n} 个关键词，总搜索量 {int(total):,}；单词搜索量 P50={p50:.0f}、P90={p90:.0f}（长尾偏斜明显）。"
    )
    insights.append(
        f"头部流量（搜索量前 {head_n} 词 ≈ 前4%）：贡献总搜索量的 {head_share:.1f}%（{int(head_vol):,}），"
        f"头部集中度{'极高' if head_share >= 60 else '较高' if head_share >= 40 else '中等'}。"
    )
    insights.append(
        f"长尾（词量后80%，共 {len(tail)} 词）合计搜索量 {int(tail_vol):,}，"
        f"占总搜索量 {100.0 * tail_vol / total:.1f}%；共性见各维长尾分布。"
    )

    # fill rate warnings
    low_dims = []
    for d in ATTR_DIMS:
        fr = meta["fill_rates"].get(d, 0)
        if fr < 5.0:
            low_dims.append(f"{d}({fr:.1f}%)")
    if low_dims:
        insights.append(
            "填充率警告（<5%）：" + "、".join(low_dims) + " —— 这些维度当前不适合作为主战场，仅作弱信号参考。"
        )
    high_dims = sorted(
        ((d, meta["fill_rates"][d]) for d in ATTR_DIMS),
        key=lambda x: -x[1],
    )[:3]
    insights.append(
        "属性填充率最高三维："
        + "、".join(f"{d}({fr:.1f}%)" for d, fr in high_dims)
        + "，更适合作为主分析轴。"
    )

    # top tags in key dims
    for dim in ("目标人群", "功效", "产品形态", "风格属性", "痛点/需求", "修饰/价值", "产地/文化"):
        arr = aggs.get(dim) or []
        if not arr:
            continue
        top = arr[0]
        if top["搜索量"] <= 0:
            continue
        insights.append(
            f"「{dim}」头部标签「{top['标签']}」：词数 {top['词数']}，分摊搜索量 {int(top['搜索量']):,}（占该维 {top['占比']:.1f}%），Top词「{top['Top词']}」。"
        )

    # head commons
    head_commons = []
    for dim in ("产品形态", "风格属性", "功效", "通用/品类", "修饰/价值"):
        dist = attr_distribution(head, dim, 3)
        if dist:
            labels = "、".join(f"{x['标签']}({int(x['搜索量']):,})" for x in dist[:3])
            head_commons.append(f"{dim}: {labels}")
    if head_commons:
        insights.append("头部词属性共性 → " + "；".join(head_commons))

    # tail commons
    tail_commons = []
    for dim in ("产品形态", "风格属性", "功效", "通用/品类"):
        dist = attr_distribution(tail, dim, 3)
        if dist:
            labels = "、".join(f"{x['标签']}(词{x['词数']})" for x in dist[:3])
            tail_commons.append(f"{dim}: {labels}")
    if tail_commons:
        insights.append("长尾词属性共性 → " + "；".join(tail_commons))

    # cooccurrence highlights
    for pair in (("目标人群", "功效"), ("产品形态", "风格属性"), ("痛点/需求", "功效")):
        co = cooccurrence(rows, pair[0], pair[1], 3)
        if co:
            bits = "、".join(f"{c['标签A']}×{c['标签B']}({int(c['搜索量']):,})" for c in co[:3])
            insights.append(f"高量交叉「{pair[0]}×{pair[1]}」Top：{bits}")

    # blue ocean
    blue = blue_ocean_window(rows, 5)
    if blue:
        bits = "、".join(f"「{b['关键词']}」({int(b['搜索量']):,}/{b['蓝海度']})" for b in blue[:5])
        insights.append(f"高需求低竞争窗口（蓝海/温和优先）：{bits}")

    # audience coverage
    aud_n = sum(1 for r in rows if r.get("目标人群"))
    insights.append(
        f"人群覆盖率：{100.0 * aud_n / n:.1f}%（{aud_n}/{n}）；"
        + ("人群信号稀疏，投放/选品勿过度依赖人群维。" if aud_n / n < 0.05 else "可结合人群维做细分投放。")
    )

    return insights


# ---------------------------------------------------------------------------
# HTML (aligned to linkfox-report-generator template-analysis.html style)
# ---------------------------------------------------------------------------

_TEMPLATE_CSS = """
:root{--color-bg:#fff;--color-surface:#fff;--color-border:#e8eaed;--color-text-primary:#1a1a2e;--color-text-secondary:#5a5a72;--color-text-muted:#8e8ea0;--color-accent:#4f46e5;--color-accent-light:#eef2ff;--sentiment-positive:#10b981;--sentiment-positive-bg:#ecfdf5;--sentiment-neutral:#f59e0b;--sentiment-neutral-bg:#fffbeb;--sentiment-negative:#ef4444;--sentiment-negative-bg:#fef2f2;--priority-high:#ef4444;--priority-medium:#f59e0b;--priority-low:#6b7280;--space-xs:4px;--space-sm:8px;--space-md:16px;--space-lg:24px;--space-xl:32px;--space-2xl:48px;--font-sans:-apple-system,BlinkMacSystemFont,"Segoe UI","PingFang SC","Hiragino Sans GB","Microsoft YaHei",sans-serif;--font-mono:"SF Mono","Fira Code","Consolas",monospace;--text-xs:12px;--text-sm:13px;--text-base:14px;--text-lg:16px;--text-xl:20px;--text-2xl:24px;--text-3xl:32px;--radius-sm:6px;--radius-md:10px;--radius-lg:14px;--shadow-sm:0 1px 3px rgba(0,0,0,.04),0 1px 2px rgba(0,0,0,.06);--shadow-md:0 4px 12px rgba(0,0,0,.06),0 2px 4px rgba(0,0,0,.04)}
*,*::before,*::after{box-sizing:border-box;margin:0;padding:0}
html{width:100%;overflow-x:hidden}
body{font-family:var(--font-sans);font-size:var(--text-base);line-height:1.6;color:var(--color-text-primary);background:var(--color-bg);-webkit-font-smoothing:antialiased;width:100%;max-width:100%;overflow-x:hidden}
.report-container{width:100%;max-width:1200px;margin:0 auto;padding:var(--space-xl) var(--space-lg);display:flex;flex-wrap:wrap;gap:var(--space-xl);align-items:flex-start}
.report-main{flex:1;min-width:0;padding-right:8px}
.report-header{background:linear-gradient(135deg,#4f46e5 0%,#7c3aed 50%,#6366f1 100%);border-radius:var(--radius-lg);padding:var(--space-xl);margin-bottom:var(--space-lg);color:#fff}
.report-header h1{font-size:clamp(18px,5vw,32px);font-weight:700;color:#fff;margin-bottom:var(--space-xs);word-wrap:break-word}
.report-header .report-subtitle{font-size:clamp(13px,2vw,16px);color:rgba(255,255,255,.8)}
.report-header .report-meta{font-size:var(--text-xs);color:rgba(255,255,255,.6);margin-top:var(--space-xs)}
.kpi-grid{display:flex;flex-wrap:wrap;gap:var(--space-sm);margin-bottom:var(--space-lg);padding:var(--space-sm) 0;border-bottom:1px solid var(--color-border)}
.kpi-card{display:flex;align-items:baseline;gap:var(--space-sm);padding:var(--space-sm) var(--space-md);flex:0 1 auto;min-width:fit-content}
.kpi-card .kpi-label{font-size:clamp(11px,2vw,13px);color:var(--color-text-muted);white-space:nowrap}
.kpi-card .kpi-value{font-size:clamp(14px,3vw,20px);font-weight:600;color:var(--color-text-primary);font-family:var(--font-mono)}
.content-section{position:relative;background:var(--color-surface);border-radius:var(--radius-md);padding:var(--space-lg) var(--space-xl);margin-bottom:var(--space-lg);box-shadow:var(--shadow-sm)}
.content-section h2{font-size:clamp(15px,3vw,20px);font-weight:600;color:var(--color-text-primary);margin-bottom:var(--space-md);padding-bottom:var(--space-sm);border-bottom:1px solid var(--color-border)}
.content-section h3{font-size:clamp(14px,2.5vw,16px);font-weight:600;color:var(--color-text-primary);margin:var(--space-lg) 0 var(--space-sm) 0}
.content-section h4{font-size:clamp(13px,2vw,14px);font-weight:600;color:var(--color-text-secondary);margin:var(--space-md) 0 var(--space-xs) 0}
.content-section p{color:var(--color-text-secondary);margin-bottom:var(--space-md);line-height:1.7}
.data-table-wrapper{overflow-x:auto;margin:var(--space-md) 0;border-radius:var(--radius-sm);-webkit-overflow-scrolling:touch}
.data-table{min-width:100%;width:max-content;border-collapse:collapse;font-size:var(--text-sm)}
.data-table thead{position:sticky;top:0;z-index:1}
.data-table th{background:var(--color-bg);font-weight:600;color:var(--color-text-secondary);padding:var(--space-sm) var(--space-md);text-align:left;border-bottom:2px solid var(--color-border);white-space:nowrap}
.data-table td{padding:var(--space-sm) var(--space-md);border-bottom:1px solid var(--color-border);color:var(--color-text-primary);vertical-align:middle;white-space:nowrap}
.data-table tbody tr:nth-child(even){background:#fafbfc}
.data-table tbody tr:hover{background:var(--color-accent-light)}
.data-table .num{text-align:right;font-family:var(--font-mono);font-size:var(--text-xs)}
.tag{display:inline-block;padding:2px 8px;border-radius:4px;font-size:var(--text-xs);font-weight:500;line-height:1.6;vertical-align:middle}
.tag-positive{background:var(--sentiment-positive-bg);color:var(--sentiment-positive)}
.tag-negative{background:var(--sentiment-negative-bg);color:var(--sentiment-negative)}
.tag-accent{background:var(--color-accent-light);color:var(--color-accent)}
.tag-muted{background:var(--color-bg);color:var(--color-text-muted)}
.insight-list{list-style:none;margin:var(--space-md) 0}
.insight-list li{padding:var(--space-sm) var(--space-md);margin-bottom:var(--space-sm);background:var(--color-bg);border-radius:var(--radius-sm);font-size:var(--text-sm);color:var(--color-text-secondary);position:relative;padding-left:var(--space-xl)}
.insight-list li::before{content:'';position:absolute;left:var(--space-md);top:50%;transform:translateY(-50%);width:6px;height:6px;border-radius:50%;background:var(--color-accent)}
.insight-list li.priority-high::before{background:var(--priority-high)}
.insight-list li.priority-medium::before{background:var(--priority-medium)}
.insight-list li.priority-low::before{background:var(--priority-low)}
.summary-box{background:var(--color-accent-light);border-left:4px solid var(--color-accent);border-radius:var(--radius-md);padding:var(--space-lg) var(--space-xl);margin:var(--space-lg) 0}
.summary-box h4{color:var(--color-accent);font-size:var(--text-lg);font-weight:600;margin-bottom:var(--space-sm)}
.summary-box p{font-size:var(--text-base);color:var(--color-text-secondary);line-height:1.7}
.chart-container{width:100%;max-width:100%;overflow-x:auto;margin:var(--space-md) auto;border-radius:var(--radius-sm);min-height:200px;text-align:center}
.chart-container canvas{display:block;margin:0 auto;max-width:100%;height:auto;width:auto}
.chart-row{display:grid;gap:var(--space-md);margin:var(--space-md) 0}
.chart-row.cols-2{grid-template-columns:repeat(2,1fr)}
@media(max-width:768px){.chart-row.cols-2{grid-template-columns:1fr}}
.report-footer{text-align:center;padding:var(--space-xl) 0 var(--space-md);font-size:var(--text-xs);color:var(--color-text-muted);border-top:1px solid var(--color-border);margin-top:var(--space-xl)}
.data-source{font-size:var(--text-xs);color:var(--color-text-muted);margin-top:var(--space-md);padding-top:var(--space-sm);border-top:1px dashed var(--color-border)}
.data-source .ds-label{font-weight:600}
.data-source .ds-tool{background:var(--color-bg);padding:1px 6px;border-radius:4px;font-family:var(--font-mono);font-size:11px;color:var(--color-text-secondary)}
.data-source .ds-time{color:var(--color-text-muted)}
.data-source .ds-computed{margin-top:var(--space-xs);font-size:11px;color:var(--color-text-muted);font-style:italic}
.toc-sidebar{width:220px;flex-shrink:0;position:fixed;right:var(--space-lg);top:var(--space-lg);height:auto;max-height:calc(100vh - 48px);background:var(--color-surface);border:1px solid var(--color-border);border-radius:var(--radius-md);padding:var(--space-md) 0;font-size:var(--text-sm);overflow-y:auto;z-index:10;max-width:0;overflow:hidden;transition:max-width .3s ease,padding .3s ease}
.toc-sidebar:hover{max-width:220px;padding:var(--space-md) 0}
.toc-sidebar::-webkit-scrollbar{width:4px}
.toc-sidebar::-webkit-scrollbar-thumb{background:#cbd5e1;border-radius:2px}
.toc-sidebar::-webkit-scrollbar-track{background:transparent}
.toc-sidebar-title{font-size:var(--text-xs);font-weight:700;text-transform:uppercase;letter-spacing:.1em;color:var(--color-text-muted);padding:0 var(--space-lg) var(--space-md)}
.toc-sidebar ul{list-style:none;margin:0;padding:0}
.toc-sidebar li a{display:block;padding:var(--space-sm) var(--space-lg);color:var(--color-text-secondary);text-decoration:none;font-size:var(--text-sm);border-radius:var(--radius-sm);margin:2px var(--space-sm);white-space:nowrap;overflow:hidden;text-overflow:ellipsis;transition:background .15s,color .15s;border-left:3px solid transparent}
.toc-sidebar li a:hover{background:var(--color-accent-light);color:var(--color-accent)}
.toc-sidebar li a.active{background:var(--color-accent-light);color:var(--color-accent);font-weight:600;border-left-color:var(--color-accent)}
.toc-sidebar .toc-sub{list-style:none;margin:0;padding:0 0 0 12px}
.toc-sidebar .toc-sub li a{font-size:var(--text-xs);padding:4px var(--space-md) 4px var(--space-lg);color:var(--color-text-muted)}
.toc-sidebar::before{content:'\\1F4C1';position:absolute;left:50%;top:50%;transform:translate(-50%,-50%);font-size:20px;opacity:.5;transition:opacity .3s ease;white-space:nowrap}
.toc-sidebar:hover::before{opacity:0}
.progress-bar-wrapper{margin:var(--space-xs) 0}
.progress-bar-label{display:flex;justify-content:space-between;font-size:var(--text-xs);color:var(--color-text-secondary);margin-bottom:2px}
.progress-bar{height:6px;background:var(--color-bg);border-radius:3px;overflow:hidden}
.progress-bar .fill{height:100%;border-radius:3px;background:var(--color-accent);transition:width .3s ease}
@media(max-width:1200px){.toc-sidebar{display:none}.report-container{display:block}}
@media(max-width:600px){.report-container{padding:var(--space-md)}}
"""

_CANVAS_JS = """
var DEFAULT_PALETTE=['#4f46e5','#06b6d4','#8b5cf6','#f59e0b','#10b981','#ef4444','#ec4899','#6366f1'];
function _normSeries(a,b){var l,d;if(Array.isArray(a)){l=a;d=b}else if(a&&typeof a==='object'){l=a.labels;d=a.datasets}d=(d||[]).map(function(s,i){var o=s||{};return{label:o.label||o.name||('系列'+(i+1)),values:o.data||o.values||[],color:o.color||(o.colors&&o.colors.length===1?o.colors[0]:DEFAULT_PALETTE[i%DEFAULT_PALETTE.length])}});return{labels:l||[],datasets:d}}
function _roundedBar(c,x,y,w,h,r){if(h<=0)return;r=Math.min(r,w/2,h);c.beginPath();c.moveTo(x,y+h);c.lineTo(x,y+r);c.quadraticCurveTo(x,y,x+r,y);c.lineTo(x+w-r,y);c.quadraticCurveTo(x+w,y,x+w,y+r);c.lineTo(x+w,y+h);c.closePath();c.fill()}
function drawBar(id,arg1,arg2){var cv=document.getElementById(id);if(!cv||!cv.getContext)return;var p=_normSeries(arg1,arg2);var labels=p.labels,datasets=p.datasets;if(labels.length===0||datasets.length===0)return;var ctx=cv.getContext('2d');var W=cv.width,H=cv.height;var pL=60,pR=24,pT=24,pB=64;var cW=W-pL-pR,cH=H-pT-pB;var n=labels.length,m=datasets.length;var gW=cW/n;var bW=Math.min(gW*.72/m,46);var mx=0;datasets.forEach(function(ds){ds.values.forEach(function(v){if(Math.abs(v)>mx)mx=Math.abs(v)})});if(mx===0)mx=1;ctx.fillStyle='#fff';ctx.fillRect(0,0,W,H);for(var i=0;i<=5;i++){var y=pT+cH-(cH*i/5);ctx.strokeStyle=i===0?'#cbd5e1':'#eef2f7';ctx.lineWidth=1;ctx.beginPath();ctx.moveTo(pL,y);ctx.lineTo(pL+cW,y);ctx.stroke();var tv=mx*i/5;var ts=tv>=10000?(tv/10000).toFixed(1)+'万':tv>=1?tv.toFixed(0):tv.toFixed(2);ctx.fillStyle='#94a3b8';ctx.font='500 11px -apple-system,sans-serif';ctx.textAlign='right';ctx.fillText(ts,pL-8,y+4)}datasets.forEach(function(ds,di){labels.forEach(function(lbl,li){var val=ds.values[li]||0;var bH=(Math.abs(val)/mx)*cH;var ox=(gW-bW*m)/2+di*bW;var x=pL+li*gW+ox;var y=pT+cH-bH;var col=ds.color;var g=ctx.createLinearGradient(x,y,x,y+bH);g.addColorStop(0,col);g.addColorStop(1,col+'cc');ctx.fillStyle=g;_roundedBar(ctx,x+1,y,bW-2,bH,4);if(m===1&&bH>16){ctx.fillStyle='#1e293b';ctx.font='600 10.5px -apple-system,sans-serif';ctx.textAlign='center';var ls=Math.abs(val)>=10000?(val/10000).toFixed(1)+'万':Math.abs(val)>=1?val.toFixed(0):(val*100).toFixed(1)+'%';ctx.fillText(ls,x+(bW-2)/2,y-5)}})});ctx.fillStyle='#475569';ctx.font='500 11px -apple-system,sans-serif';ctx.textAlign='center';labels.forEach(function(lbl,li){ctx.fillText(lbl,pL+li*gW+gW/2,pT+cH+18)})}
"""

_TOC_JS = """
document.addEventListener('DOMContentLoaded',function(){var toc=document.getElementById('toc');if(!toc)return;var ul=toc.querySelector('ul');var sections=document.querySelectorAll('.content-section');if(sections.length<2){toc.style.display='none';return}sections.forEach(function(sec,i){var h2=sec.querySelector('h2');if(!h2)return;var id='toc-s-'+i;sec.id=id;var li=document.createElement('li');var a=document.createElement('a');a.href='#'+id;a.textContent=h2.textContent.trim();li.appendChild(a);var h3s=sec.querySelectorAll('h3');if(h3s.length>0){var sub=document.createElement('ul');sub.className='toc-sub';h3s.forEach(function(h3,j){var sid='toc-s-'+i+'-'+j;h3.id=sid;var sli=document.createElement('li');var sa=document.createElement('a');sa.href='#'+sid;sa.textContent=h3.textContent.trim();sli.appendChild(sa);sub.appendChild(sli)});li.appendChild(sub)}ul.appendChild(li)});var links=toc.querySelectorAll('a[href^="#"]');var targets=[];links.forEach(function(a){var el=document.getElementById(a.getAttribute('href').slice(1));if(el)targets.push({link:a,el:el})});function onScroll(){var cur=targets[0];for(var i=0;i<targets.length;i++){if(targets[i].el.getBoundingClientRect().top<=100)cur=targets[i]}links.forEach(function(a){a.classList.remove('active')});if(cur)cur.link.classList.add('active')}window.addEventListener('scroll',onScroll,{passive:true});onScroll();toc.addEventListener('click',function(e){var a=e.target.closest('a[href^="#"]');if(!a)return;var t=document.getElementById(a.getAttribute('href').slice(1));if(t){e.preventDefault();t.scrollIntoView({behavior:'smooth',block:'start'})}})});
"""


def canvas_bar(chart_id, items, height=340):
    """Returns (html_container, js_init_code) for a Canvas bar chart."""
    if not items:
        return "<p style='font-size:13px;color:var(--color-text-muted);'>暂无数据</p>", ""
    labels = [label[:22] + ("\u2026" if len(label) > 22 else label) for label, _ in items]
    values = [int(val) for _, val in items]
    labels_js = json.dumps(labels, ensure_ascii=False)
    values_js = json.dumps(values)
    container = '<div class="chart-container"><canvas id="{}" width="1024" height="{}"></canvas></div>'.format(chart_id, height)
    js = 'drawBar("{}", {}, [{{"label":"搜索量","data":{},"color":"#4f46e5"}}]);'.format(chart_id, labels_js, values_js)
    return container, js


def table_html(headers, rows, max_rows=25):
    th = "".join(
        '<th{}>{}</th>'.format(' class="num"' if i > 0 else '', html.escape(str(h)))
        for i, h in enumerate(headers)
    )
    body = []
    for r in rows[:max_rows]:
        tds = "".join(
            '<td{}>{}</td>'.format(' class="num"' if i > 0 else '', html.escape(str(c)))
            for i, c in enumerate(r)
        )
        body.append("<tr>{}</tr>".format(tds))
    more = ""
    if len(rows) > max_rows:
        more = '<p style="font-size:12px;color:var(--color-text-muted);margin-top:8px;">仅展示前 {} 行，完整数据见 Excel。</p>'.format(max_rows)
    return '<div class="data-table-wrapper"><table class="data-table"><thead><tr>{}</tr></thead><tbody>{}</tbody></table></div>{}'.format(th, "".join(body), more)


def render_html(
    title: str,
    category: str,
    rows: List[Dict[str, Any]],
    aggs: Dict[str, List],
    insights: List[str],
    meta: Dict,
    head: List,
    tail: List,
) -> str:
    n = len(rows)
    total = meta["total_vol"]
    generated = datetime.now().strftime("%Y-%m-%d %H:%M CST")

    # Canvas chart JS collector
    canvas_scripts = []
    _cid = [0]

    def make_bar(items, height=340):
        _cid[0] += 1
        cid = "chart_cosmo_{}".format(_cid[0])
        h, js = canvas_bar(cid, items, height=height)
        if js:
            canvas_scripts.append(js)
        return h

    # KPI cards
    kpi_cards = (
        '<div class="kpi-grid">'
        '<div class="kpi-card"><div class="kpi-label">关键词数</div><div class="kpi-value">{:,}</div></div>'
        '<div class="kpi-card"><div class="kpi-label">总搜索量</div><div class="kpi-value">{:,}</div></div>'
        '<div class="kpi-card"><div class="kpi-label">人群覆盖率</div><div class="kpi-value">{:.1f}%</div></div>'
        '<div class="kpi-card"><div class="kpi-label">P50 / P90</div><div class="kpi-value">{:.0f} / {:.0f}</div></div>'
        '<div class="kpi-card"><div class="kpi-label">头部4%贡献</div><div class="kpi-value">{:.1f}%</div></div>'
        '<div class="kpi-card"><div class="kpi-label">可用维数</div><div class="kpi-value">{}</div></div>'
        '</div>'
    ).format(n, int(total), meta["audience_coverage"], meta["p50"], meta["p90"], meta["head_share"], meta["fill_dims_ok"])

    # Fill rate table
    fill_rows = [
        [d, "{:.1f}%".format(meta["fill_rates"][d]),
         '<span class="tag tag-negative">偏低</span>' if meta["fill_rates"][d] < 5
         else '<span class="tag tag-positive">可用</span>']
        for d in ATTR_DIMS
    ]
    fill_table = table_html(["维度", "填充率", "建议"], fill_rows, 20)

    # Dimension sections
    dim_sections = []
    for d in ATTR_DIMS:
        arr = aggs.get(d) or []
        fr = meta["fill_rates"][d]
        warn = ' <span class="tag tag-negative">填充率偏低</span>' if fr < 5 else ""
        bar_html = make_bar([(x["标签"], x["搜索量"]) for x in arr[:12]])
        tbl = table_html(
            ["标签", "词数", "搜索量(等分)", "中位", "占比%", "Top词"],
            [[x["标签"], x["词数"], int(x["搜索量"]), x["中位搜索量"], x["占比"], x["Top词"]] for x in arr],
            20,
        )
        dim_sections.append(
            '<h3>{} <span style="font-size:12px;font-weight:400;color:var(--color-text-muted);">填充率 {:.1f}%{}</span></h3>{}{}'.format(
                html.escape(d), fr, warn, bar_html, tbl)
        )

    # Head/tail
    def subset_attr_block(name, subset, dims):
        blocks = []
        for d in dims:
            dist = attr_distribution(subset, d, 8)
            if not dist:
                continue
            bar_html = make_bar([(x["标签"], x["搜索量"]) for x in dist[:8]], height=280)
            blocks.append("<h4>{}</h4>{}".format(html.escape(d), bar_html))
        no_data = '<p style="color:var(--color-text-muted);font-size:13px;">无属性命中</p>'
        return '<h3>{}（n={}）</h3>{}'.format(html.escape(name), len(subset), "".join(blocks) or no_data)

    head_html = subset_attr_block(
        "头部流量共性（搜索量前4%）",
        head,
        ["产品形态", "风格属性", "功效", "通用/品类", "修饰/价值", "目标人群", "痛点/需求"],
    )
    tail_html = subset_attr_block(
        "长尾共性（词量后80%）",
        tail,
        ["产品形态", "风格属性", "功效", "通用/品类"],
    )

    # Co-occurrence
    co_blocks = []
    for a, b in (("目标人群", "功效"), ("产品形态", "风格属性"), ("痛点/需求", "功效"), ("修饰/价值", "产品形态")):
        co = cooccurrence(rows, a, b, 15)
        if not co:
            continue
        tbl = table_html(
            ["标签A", "标签B", "分摊搜索量", "词数"],
            [[c["标签A"], c["标签B"], int(c["搜索量"]), c["词数"]] for c in co],
            15,
        )
        co_blocks.append("<h4>{} \u00d7 {}</h4>{}".format(html.escape(a), html.escape(b), tbl))

    # Blue ocean
    blue = blue_ocean_window(rows, 25)
    blue_tbl = table_html(
        ["关键词", "搜索量", "蓝海度", "竞品数", "推荐行动"],
        [[b["关键词"], int(b["搜索量"]), b["蓝海度"], b["竞品数"], b["推荐行动"]] for b in blue],
        25,
    )

    # Insights
    insight_lis = "".join(
        '<li class="priority-{}">{}</li>'.format(
            "high" if i < 3 else "medium" if i < 6 else "low",
            html.escape(s)
        )
        for i, s in enumerate(insights)
    )

    # --- Deep dive: Top 5 dimensions by fill rate ---
    top_dims = sorted(
        [(d, meta["fill_rates"][d]) for d in ATTR_DIMS if meta["fill_rates"][d] > 0],
        key=lambda x: -x[1]
    )[:5]

    deep_dive_blocks = []
    deep_dive_insights = []

    for rank, (td, fr) in enumerate(top_dims, 1):
        arr = aggs.get(td) or []
        if not arr:
            continue

        tag_count = len(arr)
        top3 = arr[:3]
        top3_share = sum(x["占比"] for x in top3)
        top1 = arr[0]
        is_concentrated = top1["占比"] >= 40
        vol_list = [x["搜索量"] for x in arr]
        max_vol = max(vol_list) if vol_list else 0
        med_vol = statistics.median(vol_list) if vol_list else 0

        # Build insight per dim
        concentration_label = "高度集中" if is_concentrated else "分散均衡"
        insight_text = (
            "「{}」填充率 {:.1f}%，共 {} 个标签。"
            "Top3 标签「{}」合计占比 {:.1f}%，{}。"
        ).format(
            td, fr, tag_count,
            "、".join(x["标签"] for x in top3),
            top3_share,
            "头部效应显著，适合聚焦核心标签" if is_concentrated else "标签分布较分散，可多角度覆盖"
        )
        deep_dive_insights.append(insight_text)

        # Bar chart for this dim (top 12 tags)
        bar_html = make_bar([(x["标签"], x["搜索量"]) for x in arr[:12]])

        # Progress bars for top 5 tags share
        progress_bars = []
        for x in arr[:5]:
            pct = x["占比"]
            label = x["标签"][:12] + ("\u2026" if len(x["标签"]) > 12 else x["标签"])
            progress_bars.append(
                ('<div class="progress-bar-wrapper">'
                '<div class="progress-bar-label"><span>{} ({:,})</span><span>{:.1f}%</span></div>'
                '<div class="progress-bar"><div class="fill" style="width:{}%"></div></div>'
                '</div>'
                ).format(html.escape(label), int(x["搜索量"]), pct, min(pct, 100))
            )

        # Detail table for this dim (all tags)
        detail_tbl = table_html(
            ["标签", "词数", "搜索量(等分)", "中位", "占比%", "Top词"],
            [[x["标签"], x["词数"], int(x["搜索量"]), x["中位搜索量"], x["占比"], x["Top词"]] for x in arr],
            15,
        )

        # Cross-dim co-occurrence with other top dims
        cross_blocks = []
        for other_td, _ in top_dims:
            if other_td == td:
                continue
            co = cooccurrence(rows, td, other_td, 5)
            if not co:
                continue
            co_tbl = table_html(
                ["{}标签".format(td[:4]), "{}标签".format(other_td[:4]), "分摊搜索量", "词数"],
                [[c["标签A"], c["标签B"], int(c["搜索量"]), c["词数"]] for c in co],
                5,
            )
            cross_blocks.append("<h4>{} \u00d7 {}</h4>{}".format(html.escape(td), html.escape(other_td), co_tbl))

        concentration_tag = '<span class="tag tag-accent">{}</span>'.format(concentration_label)

        deep_dive_blocks.append(
            ('<h3>{}. {} <span style="font-size:12px;font-weight:400;color:var(--color-text-muted);">'
            '填充率 {:.1f}% · {} 标签 · Top3占比 {:.1f}% {}</span></h3>'
            '<div class="summary-box"><h4>维度洞察</h4><p>{}</p></div>'
            '{}'
            '<h4>Top 5 标签占比</h4>{}'
            '<h4>全标签明细</h4>{}'
            '{}'
            ).format(
                rank, html.escape(td), fr, tag_count, top3_share, concentration_tag,
                html.escape(insight_text),
                bar_html,
                "".join(progress_bars),
                detail_tbl,
                "".join(cross_blocks) if cross_blocks else ""
            )
        )

    deep_dive_insight_lis = "".join(
        '<li class="priority-{}">{}</li>'.format(
            "high" if i < 2 else "medium" if i < 4 else "low",
            html.escape(s)
        )
        for i, s in enumerate(deep_dive_insights)
    )

    # Renumber sections: shift original 5,6 to 6,7 and insert deep dive as 5
    canvas_js_all = "\n".join(canvas_scripts)

    html_out = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>{TITLE}</title>
<style>{CSS}</style>
</head>
<body>
<div class="report-container">
<div class="report-main">

<div class="report-header">
  <h1>{TITLE}</h1>
  <div class="report-subtitle">COSMO 属性市场分析</div>
  <div class="report-meta">品类：{CATEGORY} · 生成时间：{GENERATED} · 口径：多标签搜索量等分 · 源：COSMO 属性 CSV</div>
</div>

{KPI_CARDS}

<section class="content-section">
  <h2>1. 总览与填充率</h2>
  <div class="summary-box">
    <h4>口径说明</h4>
    <p>属性多标签（如 a|b）在维度汇总时对搜索量做<strong>等分</strong>；未标注「-」不计入该维。数字均来自输入 CSV，未外推。</p>
  </div>
  {FILL_TABLE}
  <div class="data-source">
    <span class="ds-label">数据源：</span>
    <span class="ds-tool">bailing-cerebro-keyword-analysis</span>
    <span class="ds-time">· {GENERATED}</span>
    <div class="ds-computed">
      <span class="ds-label">计算指标：</span>
      填充率 = 该维有效标签词数 ÷ 总词数 · 头部贡献 = 前4%词搜索量 ÷ 总搜索量 · 多标签等分：词搜索量 ÷ 标签数
    </div>
  </div>
</section>

<section class="content-section">
  <h2>2. 数学 / 科学洞察</h2>
  <ul class="insight-list">{INSIGHT_LIS}</ul>
</section>

<section class="content-section">
  <h2>3. 头部 vs 长尾</h2>
  {HEAD_HTML}
  {TAIL_HTML}
</section>

<section class="content-section">
  <h2>4. 属性维度流量副表（COSMO 13 维全景）</h2>
  {DIM_SECTIONS}
</section>

<section class="content-section">
  <h2>5. 产品深耕分析（Top 5 维度下钻）</h2>
  <div class="summary-box">
    <h4>分析逻辑</h4>
    <p>在 COSMO 13 维全景基础上，自动选取填充率最高的前 5 个维度做下钻分析。每个维度展示标签集中度、Top 5 占比分布、全标签明细和跨维度共现，帮助识别该品类的核心属性轴。</p>
  </div>
  {DEEP_DIVE_BLOCKS}
  <h3>深耕小结</h3>
  <ul class="insight-list">{DEEP_DIVE_INSIGHT_LIS}</ul>
</section>

<section class="content-section">
  <h2>6. 维度交叉（共现）</h2>
  {CO_BLOCKS}
</section>

<section class="content-section">
  <h2>7. 蓝海 × 搜索量窗口</h2>
  <p style="color:var(--color-text-muted);font-size:13px;">优先「蓝海/温和」且仍有搜索量的词；激烈/极激烈已排除。</p>
  {BLUE_TBL}
</section>

<div class="report-footer">COSMO 市场分析 · bailing-cosmo-market-analysis · 离线自包含 HTML（无 CDN）</div>

</div>
<aside class="toc-sidebar" id="toc">
  <div class="toc-sidebar-title">目录</div>
  <ul></ul>
</aside>
</div>

<script>
{CANVAS_JS}
</script>
<script>
document.addEventListener('DOMContentLoaded', function() {{
  try {{
    {CANVAS_JS_ALL}
  }} catch(e) {{ console.error('Canvas chart init failed:', e); }}
}});
</script>
<script>
{TOC_JS}
</script>
</body>
</html>
""".format(
        TITLE=html.escape(title),
        CSS=_TEMPLATE_CSS,
        CATEGORY=html.escape(category or "-"),
        GENERATED=generated,
        KPI_CARDS=kpi_cards,
        FILL_TABLE=fill_table,
        INSIGHT_LIS=insight_lis,
        HEAD_HTML=head_html,
        TAIL_HTML=tail_html,
        DIM_SECTIONS="".join(dim_sections),
        CO_BLOCKS="".join(co_blocks) or '<p style="color:var(--color-text-muted);font-size:13px;">有效交叉不足</p>',
        BLUE_TBL=blue_tbl,
        DEEP_DIVE_BLOCKS="".join(deep_dive_blocks),
        DEEP_DIVE_INSIGHT_LIS=deep_dive_insight_lis,
        CANVAS_JS=_CANVAS_JS,
        CANVAS_JS_ALL=canvas_js_all,
        TOC_JS=_TOC_JS,
    )

    return html_out


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(color="F8FAFC", bold=True)
SECTION_FILL = PatternFill("solid", fgColor="334155")


def _style_header(ws, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _autosize(ws, max_width: int = 42):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col[:80]:
            v = "" if cell.value is None else str(cell.value)
            length = max(length, min(len(v) + 2, max_width))
        ws.column_dimensions[letter].width = max(10, length)


def write_excel(
    path: Path,
    rows: List[Dict[str, Any]],
    aggs: Dict[str, List],
    insights: List[str],
    meta: Dict,
    title: str,
    category: str,
):
    wb = Workbook()

    # --- 说明 ---
    ws = wb.active
    ws.title = "说明"
    lines = [
        ("报告标题", title),
        ("品类", category or "-"),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M CST")),
        ("词数", meta["n"]),
        ("总搜索量", int(meta["total_vol"])),
        ("人群覆盖率%", round(meta["audience_coverage"], 2)),
        ("P50 搜索量", meta["p50"]),
        ("P90 搜索量", meta["p90"]),
        ("头部4%词数", meta["head_n"]),
        ("头部贡献占比%", round(meta["head_share"], 2)),
        ("口径", "多标签（a|b）在维度汇总时搜索量等分；空/- 视为未标"),
        ("依赖", "输入为 bailing-cerebro-keyword-analysis 产出的 COSMO 属性 CSV"),
    ]
    ws.append(["字段", "值"])
    for a, b in lines:
        ws.append([a, b])
    _style_header(ws, 2)
    ws.append([])
    ws.append(["维度", "填充率%"])
    for d in ATTR_DIMS:
        ws.append([d, round(meta["fill_rates"][d], 2)])
    _autosize(ws)

    # --- 词明细 ---
    ws2 = wb.create_sheet("词明细")
    detail_headers = CORE_COLS + ATTR_DIMS
    ws2.append(detail_headers)
    for r in rows:
        row = []
        for c in CORE_COLS:
            if c == "搜索量":
                row.append(r["搜索量"])
            else:
                row.append(r.get(c, ""))
        for d in ATTR_DIMS:
            row.append(r.get(f"{d}__raw") or ("|".join(r.get(d) or []) if r.get(d) else "-"))
        ws2.append(row)
    _style_header(ws2, len(detail_headers))
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(detail_headers))}{ws2.max_row}"
    ws2.freeze_panes = "C2"
    _autosize(ws2)

    # --- 维度汇总（全维合并） ---
    ws3 = wb.create_sheet("维度汇总")
    ws3.append(["维度", "标签", "词数", "搜索量(等分)", "中位搜索量", "占比%", "Top词", "Top词搜索量"])
    for d in ATTR_DIMS:
        for x in aggs.get(d) or []:
            ws3.append([x["维度"], x["标签"], x["词数"], x["搜索量"], x["中位搜索量"], x["占比"], x["Top词"], x["Top词搜索量"]])
    _style_header(ws3, 8)
    ws3.auto_filter.ref = f"A1:H{ws3.max_row}"
    ws3.freeze_panes = "C2"
    _autosize(ws3)

    # --- 副表汇总（分段） ---
    ws4 = wb.create_sheet("副表汇总")
    ws4.append(["段落", "维度", "标签", "词数", "搜索量(等分)", "占比%", "Top词"])
    section_map = {
        "人群汇总": ["目标人群"],
        "价值修饰": ["修饰/价值"],
        "功效": ["功效"],
        "形态": ["产品形态"],
        "风格": ["风格属性"],
        "痛点": ["痛点/需求"],
        "产地": ["产地/文化"],
        "场景节日": ["场景/活动", "季节/节日", "购买场所"],
        "材质部位": ["成分/材质", "使用部位"],
        "品类": ["通用/品类"],
    }
    for sec, dims in section_map.items():
        for d in dims:
            for x in aggs.get(d) or []:
                ws4.append([sec, d, x["标签"], x["词数"], x["搜索量"], x["占比"], x["Top词"]])
    _style_header(ws4, 7)
    ws4.auto_filter.ref = f"A1:G{ws4.max_row}"
    _autosize(ws4)

    # dedicated sheets for key dims (optional clarity)
    for sheet_name, dim in [
        ("人群汇总", "目标人群"),
        ("价值修饰", "修饰/价值"),
        ("功效", "功效"),
        ("形态", "产品形态"),
        ("风格", "风格属性"),
        ("痛点", "痛点/需求"),
        ("产地", "产地/文化"),
    ]:
        w = wb.create_sheet(sheet_name)
        w.append(["标签", "词数", "搜索量(等分)", "中位搜索量", "占比%", "Top词", "Top词搜索量"])
        for x in aggs.get(dim) or []:
            w.append([x["标签"], x["词数"], x["搜索量"], x["中位搜索量"], x["占比"], x["Top词"], x["Top词搜索量"]])
        _style_header(w, 7)
        if w.max_row > 1:
            w.auto_filter.ref = f"A1:G{w.max_row}"
        _autosize(w)

    # --- 洞察 ---
    ws5 = wb.create_sheet("洞察")
    ws5.append(["#", "结论"])
    for i, s in enumerate(insights, 1):
        ws5.append([i, s])
    _style_header(ws5, 2)
    ws5.column_dimensions["A"].width = 6
    ws5.column_dimensions["B"].width = 100
    for row in ws5.iter_rows(min_row=2, max_row=ws5.max_row, min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # --- 交叉 ---
    ws6 = wb.create_sheet("交叉共现")
    ws6.append(["维A", "标签A", "维B", "标签B", "分摊搜索量", "词数"])
    for a, b in (("目标人群", "功效"), ("产品形态", "风格属性"), ("痛点/需求", "功效"), ("修饰/价值", "产品形态")):
        for c in cooccurrence(rows, a, b, 30):
            ws6.append([c["维A"], c["标签A"], c["维B"], c["标签B"], c["搜索量"], c["词数"]])
    _style_header(ws6, 6)
    if ws6.max_row > 1:
        ws6.auto_filter.ref = f"A1:F{ws6.max_row}"
    _autosize(ws6)

    # --- 蓝海窗口 ---
    ws7 = wb.create_sheet("蓝海窗口")
    ws7.append(["关键词", "搜索量", "蓝海度", "竞品数", "推荐行动"])
    for b in blue_ocean_window(rows, 50):
        ws7.append([b["关键词"], b["搜索量"], b["蓝海度"], b["竞品数"], b["推荐行动"]])
    _style_header(ws7, 5)
    if ws7.max_row > 1:
        ws7.auto_filter.ref = f"A1:E{ws7.max_row}"
    _autosize(ws7)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------------------
# Also write CSV summaries
# ---------------------------------------------------------------------------

def write_summary_csvs(out_dir: Path, aggs: Dict[str, List], insights: List[str]):
    # master dim summary
    path = out_dir / "cosmo-dim-summary.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(
            f,
            fieldnames=["维度", "标签", "词数", "搜索量", "中位搜索量", "占比", "Top词", "Top词搜索量"],
        )
        w.writeheader()
        for d in ATTR_DIMS:
            for x in aggs.get(d) or []:
                w.writerow(x)
    insight_path = out_dir / "cosmo-insights.txt"
    insight_path.write_text("\n".join(f"{i}. {s}" for i, s in enumerate(insights, 1)), encoding="utf-8")
    return path, insight_path


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def analyze(rows: List[Dict[str, Any]]) -> Tuple[Dict, Dict[str, List], List, List, List[str]]:
    n = len(rows)
    vols = sorted(r["搜索量"] for r in rows)
    total = sum(vols)
    fill_rates = {d: fill_rate(rows, d) for d in ATTR_DIMS}
    aud_n = sum(1 for r in rows if r.get("目标人群"))
    head, tail, _ = head_tail_split(rows, 4.0, 80.0)
    head_vol = sum(r["搜索量"] for r in head)
    meta = {
        "n": n,
        "total_vol": total,
        "p50": percentile(vols, 50) if vols else 0,
        "p90": percentile(vols, 90) if vols else 0,
        "fill_rates": fill_rates,
        "audience_coverage": 100.0 * aud_n / n if n else 0,
        "head_n": len(head),
        "head_share": 100.0 * head_vol / total if total else 0,
        "fill_dims_ok": sum(1 for d in ATTR_DIMS if fill_rates[d] >= 5.0),
    }
    aggs = {d: dim_aggregate(rows, d) for d in ATTR_DIMS}
    insights = build_insights(rows, aggs, meta)
    return meta, aggs, head, tail, insights


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="COSMO 市场分析：CSV → HTML + Excel")
    ap.add_argument("--input", "-i", required=True, help="COSMO 属性 CSV 路径")
    ap.add_argument("--output-dir", "-o", required=True, help="输出目录")
    ap.add_argument("--title", default="COSMO 市场分析报告")
    ap.add_argument("--category", default="")
    args = ap.parse_args(argv)

    inp = Path(args.input).expanduser().resolve()
    out_dir = Path(args.output_dir).expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    if not inp.exists():
        raise SystemExit(f"输入不存在: {inp}")

    rows, colmap, _ = load_rows(inp)
    if not rows:
        raise SystemExit("CSV 无有效关键词行")

    meta, aggs, head, tail, insights = analyze(rows)

    slug = re.sub(r"[^\w\-]+", "-", (args.category or "cosmo").strip())[:40] or "cosmo"
    html_path = out_dir / f"cosmo-market-report-{slug}.html"
    xlsx_path = out_dir / f"cosmo-market-workbook-{slug}.xlsx"
    # also a filtered-friendly keyword csv copy summary is in workbook; export dim csv
    kw_csv = out_dir / f"cosmo-keyword-detail-{slug}.csv"

    html_body = render_html(args.title, args.category, rows, aggs, insights, meta, head, tail)
    html_path.write_text(html_body, encoding="utf-8")

    write_excel(xlsx_path, rows, aggs, insights, meta, args.title, args.category)

    # keyword detail csv
    with kw_csv.open("w", encoding="utf-8-sig", newline="") as f:
        headers = CORE_COLS + ATTR_DIMS
        w = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            row = {c: r.get(c, "") for c in CORE_COLS}
            for d in ATTR_DIMS:
                row[d] = r.get(f"{d}__raw") or ("|".join(r.get(d) or []) if r.get(d) else "-")
            w.writerow(row)

    dim_csv, insight_txt = write_summary_csvs(out_dir, aggs, insights)

    # print summary for agent
    print("=== COSMO Market Analysis Done ===")
    print(f"input: {inp}")
    print(f"keywords: {meta['n']}")
    print(f"total_volume: {int(meta['total_vol'])}")
    print(f"P50/P90: {meta['p50']:.0f}/{meta['p90']:.0f}")
    print(f"head4%_share: {meta['head_share']:.2f}%")
    print(f"audience_coverage: {meta['audience_coverage']:.2f}%")
    print(f"HTML: {html_path} ({html_path.stat().st_size} bytes)")
    print(f"XLSX: {xlsx_path} ({xlsx_path.stat().st_size} bytes)")
    print(f"CSV_detail: {kw_csv}")
    print(f"CSV_dim: {dim_csv}")
    print(f"insights: {insight_txt}")
    print("--- Top insights ---")
    for i, s in enumerate(insights[:8], 1):
        print(f"{i}. {s}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
